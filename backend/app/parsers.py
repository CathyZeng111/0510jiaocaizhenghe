from __future__ import annotations

import re
import tempfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import fitz
from docx import Document
from fastapi import UploadFile

from .schemas import Chapter, Textbook, UploadResult
from .textbook_store import cache_upload_result

SUPPORTED_EXTENSIONS = {".pdf", ".md", ".markdown", ".txt", ".docx"}
CHAPTER_RE = re.compile(
    r"^\s*((第\s*[一二三四五六七八九十百千万零〇两\d]+\s*[章节篇编部讲])|"
    r"(Chapter\s+\d+))",
    re.IGNORECASE,
)
MARKDOWN_HEADING_RE = re.compile(r"^(#{1,3})\s+(.+?)\s*$")
MAX_PREVIEW_TITLE_LEN = 80
BAD_PDF_CHARS_RE = re.compile(r"[\ufffd\x08]+")
INVISIBLE_PDF_CHARS_RE = re.compile(r"[\ufeff\u200b\u200c\u200d]")
TOC_ENTRY_RE = re.compile(r"[\ufffd\x08]{3,}.*\d+\s*$")


@dataclass
class PdfLine:
    text: str
    size: float
    bold: bool


async def parse_upload(file: UploadFile, index: int) -> UploadResult:
    filename = file.filename or f"upload_{index}"
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        return UploadResult(
            filename=filename,
            format=extension.lstrip(".") or "unknown",
            size=0,
            status="failed",
            error=f"不支持的文件格式：{extension or 'unknown'}",
        )

    textbook_id = f"book_{index:02d}"

    temp_path: Optional[Path] = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=extension) as tmp:
            temp_path = Path(tmp.name)
            size = 0
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                tmp.write(chunk)

        textbook = parse_file_path(temp_path, filename, textbook_id)

        return cache_upload_result(
            UploadResult(
                filename=filename,
                format=extension.lstrip("."),
                size=size,
                status="completed",
                textbook=textbook,
            )
        )
    except Exception as exc:
        return UploadResult(
            filename=filename,
            format=extension.lstrip("."),
            size=temp_path.stat().st_size if temp_path and temp_path.exists() else 0,
            status="failed",
            error=str(exc),
        )
    finally:
        if temp_path and temp_path.exists():
            temp_path.unlink(missing_ok=True)


def parse_file_path(path: Path, filename: Optional[str] = None, textbook_id: str = "book_01") -> Textbook:
    filename = filename or path.name
    extension = path.suffix.lower()

    if extension == ".pdf":
        return parse_pdf_path(path, filename, textbook_id)
    if extension in {".md", ".markdown"}:
        return parse_markdown(path.read_bytes(), filename, textbook_id)
    if extension == ".txt":
        return parse_txt(path.read_bytes(), filename, textbook_id)
    if extension == ".docx":
        return parse_docx(path.read_bytes(), filename, textbook_id)
    raise ValueError(f"不支持的文件格式：{extension or 'unknown'}")


def parse_pdf(raw: bytes, filename: str, textbook_id: str) -> Textbook:
    with tempfile.NamedTemporaryFile(suffix=".pdf") as tmp:
        tmp.write(raw)
        tmp.flush()
        return parse_pdf_path(Path(tmp.name), filename, textbook_id)

def parse_pdf_path(path: Path, filename: str, textbook_id: str) -> Textbook:
        document = fitz.open(path)
        page_count = document.page_count
        repeated_headers = detect_repeated_margin_text(document)
        chapters: list[Chapter] = []
        current_title = "正文"
        current_start_page = 1
        current_parts: list[str] = []
        total_chars = 0

        for page_index in range(document.page_count):
            page = document.load_page(page_index)
            page_number = page_index + 1
            page_lines = extract_pdf_text_lines(page, repeated_headers)
            if is_probable_toc_page(page_lines):
                continue

            line_index = 0
            while line_index < len(page_lines):
                line = page_lines[line_index]
                title, skip_next = detect_pdf_chapter_title_at(page_lines, line_index)
                if title:
                    if clean_title(title) == clean_title(current_title):
                        current_parts.append(line.text)
                        if skip_next:
                            current_parts.append(page_lines[line_index + 1].text)
                        line_index += 2 if skip_next else 1
                        continue
                    if current_parts:
                        chapters.append(
                            make_chapter(
                                len(chapters) + 1,
                                current_title,
                                current_start_page,
                                page_number if page_number == current_start_page else page_number - 1,
                                "\n".join(current_parts),
                            )
                        )
                    current_title = title
                    current_start_page = page_number
                    current_parts = [line.text]
                    if skip_next:
                        current_parts.append(page_lines[line_index + 1].text)
                else:
                    current_parts.append(line.text)

                line_index += 2 if skip_next else 1

            total_chars += sum(len(line.text) for line in page_lines)

        if current_parts:
            chapters.append(
                make_chapter(
                    len(chapters) + 1,
                    current_title,
                    current_start_page,
                    page_count,
                    "\n".join(current_parts),
                )
            )

        document.close()

        if not chapters:
            chapters = [make_chapter(1, "正文", 1, 1, "")]

        return Textbook(
            textbook_id=textbook_id,
            filename=filename,
            title=infer_title(filename),
            total_pages=page_count,
            total_chars=total_chars,
            chapters=chapters,
        )


def detect_repeated_margin_text(document: fitz.Document) -> set[str]:
    margin_lines: Counter[str] = Counter()
    sampled_pages = min(document.page_count, 40)

    for page_index in range(sampled_pages):
        page = document.load_page(page_index)
        blocks = page.get_text("dict").get("blocks", [])
        page_height = page.rect.height

        for block in blocks:
            if block.get("type") != 0:
                continue

            y0 = block.get("bbox", [0, 0, 0, 0])[1]
            y1 = block.get("bbox", [0, 0, 0, 0])[3]
            if y0 > page_height * 0.12 and y1 < page_height * 0.88:
                continue

            text = normalize_space(text_from_block(block))
            if 0 < len(text) <= 80:
                margin_lines[text] += 1

    threshold = max(3, sampled_pages // 4)
    return {line for line, count in margin_lines.items() if count >= threshold}


def extract_pdf_text_lines(page: fitz.Page, repeated_headers: set[str]) -> list[PdfLine]:
    lines: list[PdfLine] = []
    blocks = page.get_text("dict").get("blocks", [])
    page_height = page.rect.height

    for block in blocks:
        if block.get("type") != 0:
            continue

        y0, y1 = block.get("bbox", [0, 0, 0, 0])[1], block.get("bbox", [0, 0, 0, 0])[3]
        block_text = normalize_space(text_from_block(block))

        if not block_text:
            continue
        if block_text in repeated_headers:
            continue
        is_margin_page_number = bool(
            re.fullmatch(r"(第\s*)?\d+\s*(页)?|Page\s+\d+", block_text, re.IGNORECASE)
        )
        if (y0 < page_height * 0.08 or y1 > page_height * 0.93) and is_margin_page_number:
            continue

        lines.extend(line_candidates_from_block(block, repeated_headers))

    return lines


def line_candidates_from_block(block: dict, repeated_headers: set[str]) -> list[PdfLine]:
    candidates: list[PdfLine] = []
    for line in block.get("lines", []):
        spans = line.get("spans", [])
        raw_text = "".join(span.get("text", "") for span in spans)
        if not raw_text.strip():
            continue

        max_size = max((float(span.get("size", 0)) for span in spans), default=0)
        is_bold = any("bold" in str(span.get("font", "")).lower() for span in spans)

        for split_line in raw_text.splitlines():
            normalized = clean_pdf_text(split_line)
            if normalized and normalized not in repeated_headers:
                candidates.append(PdfLine(text=normalized, size=max_size, bold=is_bold))

    return candidates


def text_from_block(block: dict) -> str:
    lines: list[str] = []
    for line in block.get("lines", []):
        spans = line.get("spans", [])
        line_text = "".join(span.get("text", "") for span in spans)
        for split_line in line_text.splitlines():
            if split_line.strip():
                lines.append(split_line)
    return "\n".join(lines)


def parse_markdown(raw: bytes, filename: str, textbook_id: str) -> Textbook:
    text = decode_text(raw)
    chapters = chapters_from_markdown(text)
    return Textbook(
        textbook_id=textbook_id,
        filename=filename,
        title=infer_title(filename),
        total_pages=None,
        total_chars=len(text),
        chapters=chapters,
    )


def parse_txt(raw: bytes, filename: str, textbook_id: str) -> Textbook:
    text = decode_text(raw)
    chapters = chapters_from_plain_text(text)
    return Textbook(
        textbook_id=textbook_id,
        filename=filename,
        title=infer_title(filename),
        total_pages=None,
        total_chars=len(text),
        chapters=chapters,
    )


def parse_docx(raw: bytes, filename: str, textbook_id: str) -> Textbook:
    with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
        tmp.write(raw)
        tmp.flush()
        document = Document(tmp.name)

    chapters: list[Chapter] = []
    current_title = "正文"
    current_parts: list[str] = []

    for paragraph in document.paragraphs:
        text = normalize_space(paragraph.text)
        if not text:
            continue

        is_heading = paragraph.style and paragraph.style.name.lower().startswith("heading")
        title = text if is_heading else detect_chapter_title(text)

        if title:
            if current_parts:
                chapters.append(make_chapter(len(chapters) + 1, current_title, None, None, "\n".join(current_parts)))
            current_title = clean_title(title)
            current_parts = [text]
        else:
            current_parts.append(text)

    if current_parts:
        chapters.append(make_chapter(len(chapters) + 1, current_title, None, None, "\n".join(current_parts)))

    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    return Textbook(
        textbook_id=textbook_id,
        filename=filename,
        title=infer_title(filename),
        total_pages=None,
        total_chars=len(full_text),
        chapters=chapters or [make_chapter(1, "正文", None, None, full_text)],
    )


def chapters_from_markdown(text: str) -> list[Chapter]:
    chapters: list[Chapter] = []
    current_title = "正文"
    current_parts: list[str] = []

    for line in text.splitlines():
        match = MARKDOWN_HEADING_RE.match(line)
        if match:
            if current_parts:
                chapters.append(make_chapter(len(chapters) + 1, current_title, None, None, "\n".join(current_parts)))
            current_title = clean_title(match.group(2))
            current_parts = [line]
        else:
            current_parts.append(line)

    if current_parts:
        chapters.append(make_chapter(len(chapters) + 1, current_title, None, None, "\n".join(current_parts)))

    return chapters or [make_chapter(1, "正文", None, None, text)]


def chapters_from_plain_text(text: str) -> list[Chapter]:
    chapters: list[Chapter] = []
    current_title = "正文"
    current_parts: list[str] = []

    for line in text.splitlines():
        normalized = normalize_space(line)
        title = detect_chapter_title(normalized)
        if title:
            if current_parts:
                chapters.append(make_chapter(len(chapters) + 1, current_title, None, None, "\n".join(current_parts)))
            current_title = title
            current_parts = [line]
        else:
            current_parts.append(line)

    if current_parts:
        chapters.append(make_chapter(len(chapters) + 1, current_title, None, None, "\n".join(current_parts)))

    return chapters or [make_chapter(1, "正文", None, None, text)]


def detect_chapter_title(line: str) -> Optional[str]:
    normalized = normalize_space(line)
    if not normalized or len(normalized) > MAX_PREVIEW_TITLE_LEN:
        return None
    if CHAPTER_RE.match(normalized):
        return clean_title(normalized)
    return None


def detect_pdf_chapter_title(line: PdfLine) -> Optional[str]:
    if has_too_many_bad_pdf_chars(line.text):
        return None

    title = detect_chapter_title(line.text)
    if title:
        return title

    if re.fullmatch(r"\d+", line.text):
        return None

    if len(line.text) <= 50 and line.size >= 20 and not line.text.endswith(("。", ".", "；", ";", "，", ",")):
        return clean_title(line.text)

    if len(line.text) <= 50 and line.size >= 15 and line.bold and not line.text.endswith(("。", ".", "；", ";", "，", ",")):
        return clean_title(line.text)

    return None


def detect_pdf_chapter_title_at(lines: list[PdfLine], index: int) -> tuple[Optional[str], bool]:
    line = lines[index]
    title = detect_pdf_chapter_title(line)
    if title:
        if index + 1 < len(lines) and should_join_title_continuation(title, lines[index + 1]):
            return clean_title(f"{title}{lines[index + 1].text}"), True
        return title, False

    if index + 1 >= len(lines) or not re.fullmatch(r"\d+", line.text):
        return None, False

    next_line = lines[index + 1]
    next_text = next_line.text
    if (
        len(next_text) <= 40
        and not re.fullmatch(r"\d+", next_text)
        and not next_text.endswith(("。", ".", "；", ";", "，", ","))
        and (line.size >= 15 or next_line.size >= 15 or line.bold or next_line.bold)
    ):
        return clean_title(f"{line.text} {next_text}"), True

    return None, False


def should_join_title_continuation(title: str, next_line: PdfLine) -> bool:
    next_text = next_line.text.strip()
    if not next_text or len(next_text) > 4:
        return False
    if re.search(r"[。；;，,：:!?？]$", next_text):
        return False
    if re.fullmatch(r"\d+", next_text):
        return False
    if not re.search(r"[\u4e00-\u9fff]", next_text):
        return False
    if re.match(r"^(第[一二三四五六七八九十百千万零〇两\d]+章|绪)$", title):
        return True
    if re.match(r"^第[一二三四五六七八九十百千万零〇两\d]+章\s*[\u4e00-\u9fff]{1,3}$", title):
        return True
    if re.match(r"^第[一二三四五六七八九十百千万零〇两\d]+节.*[\u4e00-\u9fff]$", title) and len(title) <= 10:
        return True
    return False


def is_probable_toc_page(lines: list[PdfLine]) -> bool:
    if not lines:
        return False

    texts = [line.text for line in lines]
    joined_head = " ".join(texts[:8])
    if "目录" in joined_head:
        return True

    toc_entry_count = sum(1 for text in texts if TOC_ENTRY_RE.search(text))
    numbered_tail_count = sum(1 for text in texts if re.search(r"\d+\s*$", text) and len(text) <= 80)
    bad_line_count = sum(1 for text in texts if has_too_many_bad_pdf_chars(text))
    section_entry_count = sum(
        1
        for text in texts
        if re.match(r"^\s*(第[一二三四五六七八九十百千万零〇两\d]+[章节]|[一二三四五六七八九十]+、)", text)
        and re.search(r"\d+\s*$", text)
    )

    return (
        toc_entry_count >= 4
        or (bad_line_count >= 6 and numbered_tail_count >= 6)
        or (section_entry_count >= 8 and numbered_tail_count >= 8)
    )


def make_chapter(
    index: int,
    title: str,
    page_start: Optional[int],
    page_end: Optional[int],
    content: str,
) -> Chapter:
    stripped_content = content.strip()
    return Chapter(
        chapter_id=f"ch_{index:02d}",
        title=clean_title(title),
        page_start=page_start,
        page_end=page_end,
        content=stripped_content,
        char_count=len(stripped_content),
    )


def normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def clean_title(value: str) -> str:
    value = clean_pdf_text(value)
    value = re.sub(r"[\ufffd�]+", "", value)
    value = re.sub(r"\s+\d+\s*$", "", value)
    return normalize_space(value).strip("#").strip()


def clean_pdf_text(value: str) -> str:
    value = INVISIBLE_PDF_CHARS_RE.sub("", value)
    value = BAD_PDF_CHARS_RE.sub(" ", value)
    value = value.replace("\u2003", " ").replace("\u2002", " ").replace("\u200a", " ")
    return normalize_space(value)


def has_too_many_bad_pdf_chars(value: str) -> bool:
    bad_count = len(BAD_PDF_CHARS_RE.findall(value))
    if bad_count == 0:
        return False
    return bad_count >= 6 or bad_count / max(len(value), 1) > 0.18


def infer_title(filename: str) -> str:
    return Path(filename).stem


def decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")
